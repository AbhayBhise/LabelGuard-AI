"""Report generator — legal PDF and DOCX compliance reports."""
import io
from datetime import datetime

import numpy as np

from app.services.datatypes import ComplianceResult, ExtractedFields


class ReportGenerator:
    def generate_pdf(
        self,
        scan_id: str,
        product: dict,
        fields: ExtractedFields,
        result: ComplianceResult,
        officer: dict = None,
        annotated_image: bytes = None,
    ) -> bytes:
        """Build a ReportLab PDF with header, summary, violation table and
        annotated evidence image."""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Image,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        officer = officer or {}
        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=A4,
            rightMargin=18 * mm, leftMargin=18 * mm,
            topMargin=18 * mm, bottomMargin=18 * mm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "TitleX", parent=styles["Title"], fontSize=20,
            textColor=colors.HexColor("#1B4FD8"),
        )
        styles.add(title_style)

        story = []
        story.append(Paragraph("LabelGuard AI — Compliance Report", title_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph(
            f"Inspection under Legal Metrology (Packaged Commodities) Rules, 2011"
            f"<br/>Scan ID: {scan_id} | Date: {datetime.now().strftime('%d %b %Y %H:%M')}"
            f"<br/>Officer: {officer.get('name', '—')} ({officer.get('role', '—')})",
            styles["Normal"],
        ))
        story.append(Spacer(1, 10))

        # Product identity
        story.append(Paragraph("1. Product Identification", styles["Heading2"]))
        pdata = [
            ["Product Name", product.get("product_name", "—")],
            ["Barcode", product.get("barcode", "—")],
            ["Brand", product.get("brand", "—")],
            ["Category", product.get("category", "—")],
        ]
        pt = Table(pdata, colWidths=[40 * mm, 120 * mm])
        pt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2FF")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(pt)
        story.append(Spacer(1, 10))

        # Compliance summary
        story.append(Paragraph("2. Compliance Summary", styles["Heading2"]))
        color_map = {
            "COMPLIANT": "16A34A",
            "PARTIAL": "D97706",
            "NON_COMPLIANT": "DC2626",
        }
        status = result.status
        status_hex = color_map.get(status, "0F172A")
        story.append(Paragraph(
            f"Overall Status: <font color='#{status_hex}'>"
            f"<b>{status}</b></font>",
            styles["Normal"],
        ))
        story.append(Paragraph(
            f"Compliance Score: {result.score:.2f} / 1.00",
            styles["Normal"],
        ))
        story.append(Paragraph(
            f"Violations: {len(result.violations)}",
            styles["Normal"],
        ))
        story.append(Spacer(1, 10))

        # Violation table
        story.append(Paragraph("3. Violations", styles["Heading2"]))
        vdata = [["Rule", "Type", "Severity", "Finding"]]
        for v in result.violations:
            vdata.append([
                v.rule_id, v.violation_type, v.severity,
                (v.finding or "")[:120],
            ])
        if len(vdata) == 1:
            vdata.append(["—", "—", "—", "No violations found"])
        vt = Table(vdata, colWidths=[25 * mm, 25 * mm, 25 * mm, 85 * mm])
        vt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B4FD8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(vt)
        story.append(Spacer(1, 10))

        # Extracted fields
        story.append(Paragraph("4. Extracted Fields", styles["Heading2"]))
        fdata = [["Field", "Extracted Value"]]
        for k, v in fields.as_dict().items():
            if k == "all_text":
                continue
            fdata.append([k, str(v)[:100]])
        if len(fdata) == 1:
            fdata.append(["—", "No fields extracted"])
        ft = Table(fdata, colWidths=[60 * mm, 100 * mm])
        ft.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B4FD8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(ft)
        story.append(Spacer(1, 10))

        # Annotated image
        story.append(Paragraph("5. Annotated Label Image", styles["Heading2"]))
        if annotated_image:
            img_path = _ImageTemp(annotated_image)
            try:
                story.append(Image(
                    img_path.path,
                    width=120 * mm,
                    height=120 * mm,
                    kind="proportional",
                ))
            except Exception:
                story.append(Paragraph("(image unavailable)", styles["Normal"]))
            story.append(Spacer(1, 10))

        story.append(Paragraph(
            "6. Officer Declaration", styles["Heading2"]
        ))
        story.append(Spacer(1, 20))
        story.append(Paragraph(
            "Signature: ______________________ &nbsp;&nbsp; Date: ___________",
            styles["Normal"],
        ))

        doc.build(story)
        return buf.getvalue()

    def generate_docx(
        self,
        scan_id: str,
        product: dict,
        fields: ExtractedFields,
        result: ComplianceResult,
        officer: dict = None,
    ) -> bytes:
        from docx import Document
        from docx.shared import Pt

        officer = officer or {}
        doc = Document()
        doc.add_heading("LabelGuard AI — Compliance Report", level=0)
        p = doc.add_paragraph()
        p.add_run(
            f"Scan ID: {scan_id} | Date: {datetime.now().strftime('%d %b %Y %H:%M')}\n"
            f"Officer: {officer.get('name', '—')} ({officer.get('role', '—')})"
        )
        doc.add_heading("Product Identification", level=2)
        tab = doc.add_table(rows=1, cols=2)
        tab.style = "Light Grid Accent 1"
        hdr = tab.rows[0].cells
        hdr[0].text = "Field"
        hdr[1].text = "Value"
        for k, v in [
            ("Product Name", product.get("product_name", "—")),
            ("Barcode", product.get("barcode", "—")),
            ("Brand", product.get("brand", "—")),
            ("Category", product.get("category", "—")),
        ]:
            r = tab.add_row().cells
            r[0].text = k
            r[1].text = str(v)

        doc.add_heading(
            f"Compliance Summary: {result.status} (Score {result.score:.2f})",
            level=2,
        )
        doc.add_heading("Violations", level=2)
        vtab = doc.add_table(rows=1, cols=4)
        vtab.style = "Light Grid Accent 1"
        vh = vtab.rows[0].cells
        vh[0].text, vh[1].text, vh[2].text, vh[3].text = (
            "Rule", "Type", "Severity", "Finding",
        )
        for v in result.violations:
            r = vtab.add_row().cells
            r[0].text, r[1].text, r[2].text, r[3].text = (
                v.rule_id, v.violation_type, v.severity, v.finding or "",
            )

        doc.add_heading("Extracted Fields", level=2)
        etab = doc.add_table(rows=1, cols=2)
        etab.style = "Light Grid Accent 1"
        eh = etab.rows[0].cells
        eh[0].text, eh[1].text = "Field", "Value"
        for k, v in fields.as_dict().items():
            if k == "all_text":
                continue
            r = etab.add_row().cells
            r[0].text, r[1].text = k, str(v)

        doc.add_heading("Officer Declaration", level=2)
        doc.add_paragraph("Signature: ____________  Date: __________")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def annotate_image(self, image, violations, fields) -> np.ndarray:
        """Draw colored bounding boxes on the image."""
        import cv2

        img = image.copy()
        for v in violations:
            if not v.evidence_bbox:
                continue
            b = v.evidence_bbox
            x1, y1 = int(b["x"]), int(b["y"])
            x2, y2 = int(b["x"] + b["w"]), int(b["y"] + b["h"])
            color = (0, 0, 255) if v.severity == "CRITICAL" else (0, 140, 255)
            cv2.rectangle(img, (x1, y1), (x2, y2), color, 2)
            cv2.putText(
                img, f"{v.rule_id} {v.severity}", (x1, y1 - 6),
                cv2.FONT_HERSHEY_SIMPLEX, 0.5, color, 1, cv2.LINE_AA,
            )
        return img


class _ImageTemp:
    """Write annotated image bytes to a temp file so ReportLab can render it."""

    def __init__(self, data: bytes):
        self.path = self._write(data)

    def _write(self, data: bytes) -> str:
        import tempfile

        fd, path = tempfile.mkstemp(suffix=".jpg")
        import os

        with os.fdopen(fd, "wb") as f:
            f.write(data)
        return path
